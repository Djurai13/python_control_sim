from pathlib import Path

from openpyxl import load_workbook


class ExcelImporter:

    def import_data(
        self,
        source: str,
    ) -> dict:

        if not source.strip():
            raise ValueError(
                "Source is required."
            )

        file_path = Path(
            source
        )

        if not file_path.exists():
            raise ValueError(
                "Excel file does not exist."
            )

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        worksheet = (
            workbook.active
        )

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            return {
                "records": [],
                "row_count": 0,
                "column_count": 0,
            }

        headers = [
            str(value)
            for value in rows[0]
        ]

        records = []

        for row in rows[1:]:

            record = {}

            for header, value in zip(
                headers,
                row,
            ):
                record[
                    header
                ] = value

            records.append(
                record
            )

        return {
            "records": records,
            "row_count": len(
                records
            ),
            "column_count": len(
                headers
            ),
        }